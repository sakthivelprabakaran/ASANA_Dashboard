"""
BRD Writer - Writes Average values from Asana CSV back to BRD Excel file.
Uses win32com (COM automation) for lossless Excel file editing on Windows.
Falls back to openpyxl if win32com is not available.
"""

import os
import shutil
import logging
from typing import Dict, List, Any, Tuple
from core.brd_matcher import BrdMatcher

logger = logging.getLogger('AsanaGenerator.BrdWriter')

# Try to import win32com for lossless Excel editing
try:
    import win32com.client
    HAS_WIN32COM = True
    logger.info("win32com available — using lossless Excel editing")
except ImportError:
    HAS_WIN32COM = False
    logger.info("win32com not available — falling back to openpyxl")


class BrdWriter:
    """
    Writes performance values (Average) from Asana CSV data back into the BRD Excel file.
    Matches scenarios using normalized name matching (same as BrdMatcher).
    
    Uses win32com (COM automation) for lossless file editing when available.
    Falls back to openpyxl otherwise.
    """

    def __init__(self):
        self.matcher = BrdMatcher()

    def write_averages_to_brd(self, brd_file_path: str, sheet_name: str,
                               target_col_index: int, scenario_col_index: int,
                               csv_tasks: List[Dict],
                               output_path: str = None,
                               dry_run: bool = False) -> Dict[str, Any]:
        """
        Write Average values from CSV tasks to the BRD Excel file.
        
        Strategy: Always use openpyxl for matching (consistent row lookup).
        For actual writing: use win32com if available (lossless), else openpyxl.
        """
        if output_path is None:
            output_path = brd_file_path

        logger.info(f"BRD Writer: HAS_WIN32COM={HAS_WIN32COM}, dry_run={dry_run}")
        
        # STEP 1: Always use openpyxl for matching (consistent between preview and write)
        import openpyxl as _openpyxl
        wb = _openpyxl.load_workbook(brd_file_path, data_only=True, read_only=True)
        ws = wb[sheet_name]
        
        brd_rows = []
        for row in ws.iter_rows(min_row=1, values_only=True):
            brd_rows.append(list(row))
        wb.close()
        
        logger.info(f"Matching: BRD sheet '{sheet_name}': {len(brd_rows)} rows")
        
        scenario_to_row = self._build_scenario_lookup(brd_rows, scenario_col_index)
        logger.info(f"Matching: Built BRD lookup: {len(scenario_to_row)} scenarios")
        
        matched, unmatched = self._match_tasks(csv_tasks, scenario_to_row)
        
        # Read old values using openpyxl (data_only to get calculated values)
        for m in matched:
            row_idx = m['brd_row'] - 1  # 0-based for brd_rows list
            if row_idx < len(brd_rows) and target_col_index < len(brd_rows[row_idx]):
                old_val = brd_rows[row_idx][target_col_index]
                m['old_value'] = str(old_val) if old_val is not None else '-'
        
        logger.info(f"Matching results: {len(matched)} matched, {len(unmatched)} unmatched")
        
        # STEP 2: If dry run, return results without writing
        if dry_run:
            return {
                'matched': matched,
                'unmatched': unmatched,
                'total_written': 0,
                'total_skipped': len(unmatched),
                'output_file': output_path,
                'dry_run': True,
            }
        
        # STEP 3: Write using win32com (lossless) or openpyxl (fallback)
        if HAS_WIN32COM:
            return self._write_matched_with_win32com(
                brd_file_path, sheet_name, target_col_index,
                matched, unmatched, output_path
            )
        else:
            return self._write_matched_with_openpyxl(
                brd_file_path, sheet_name, target_col_index,
                matched, unmatched, output_path
            )

    def _build_scenario_lookup(self, brd_rows, scenario_col_index):
        """
        Build normalized scenario → list of row indices mapping from BRD data.
        Supports multiple BRD rows with the same scenario name (matched in order).
        """
        scenario_to_rows = {}  # scenario → [row1, row2, ...] (multiple occurrences)
        start_row = 3  # Skip first 2 header rows
        
        for row_idx in range(start_row - 1, len(brd_rows)):
            row = brd_rows[row_idx]
            if scenario_col_index < len(row):
                scenario_raw = row[scenario_col_index]
                normalized = self.matcher.normalize_scenario_name(scenario_raw)
                if normalized:
                    if normalized not in scenario_to_rows:
                        scenario_to_rows[normalized] = []
                    scenario_to_rows[normalized].append(row_idx + 1)  # 1-based
        
        return scenario_to_rows

    def _match_tasks(self, csv_tasks, scenario_to_rows):
        """
        Match CSV tasks to BRD rows by normalized scenario name.
        Supports multiple occurrences: 1st CSV match → 1st BRD row, 2nd → 2nd, etc.
        """
        matched = []
        unmatched = []
        
        # Track how many times each scenario has been matched (for sequential matching)
        scenario_match_count = {}
        
        for task in csv_tasks:
            scenario_name = task.get('Name', '')
            average = task.get('Average', '')
            perf_brd = task.get('Perf_BRD', '')
            prev_value = task.get('Previous Value', '')
            parent = task.get('Parent task', '')
            
            if not average or average in ['', '-', '0']:
                continue
            
            normalized = self.matcher.normalize_scenario_name(scenario_name)
            
            if normalized in scenario_to_rows:
                brd_rows_list = scenario_to_rows[normalized]
                
                # Get the occurrence count for this scenario
                occurrence = scenario_match_count.get(normalized, 0)
                
                if occurrence < len(brd_rows_list):
                    # Match to the Nth BRD row for this scenario
                    row_num = brd_rows_list[occurrence]
                    scenario_match_count[normalized] = occurrence + 1
                    
                    matched.append({
                        'scenario': scenario_name,
                        'parent': parent,
                        'brd_row': row_num,
                        'old_value': '-',
                        'new_value': average,
                        'perf_brd': perf_brd,
                        'prev_value': prev_value,
                    })
                else:
                    # More CSV occurrences than BRD rows — unmatched
                    logger.warning(f"Scenario '{scenario_name}' has more CSV occurrences "
                                 f"({occurrence + 1}) than BRD rows ({len(brd_rows_list)})")
                    unmatched.append({
                        'scenario': scenario_name,
                        'parent': parent,
                        'average': average,
                        'normalized': normalized,
                    })
            else:
                unmatched.append({
                    'scenario': scenario_name,
                    'parent': parent,
                    'average': average,
                    'normalized': normalized,
                })
        
        return matched, unmatched

    def _write_matched_with_win32com(self, brd_file_path, sheet_name, target_col_index,
                                      matched, unmatched, output_path):
        """Write pre-matched results using win32com — lossless Excel editing."""
        import pythoncom
        pythoncom.CoInitialize()
        
        excel = None
        wb = None
        
        try:
            abs_src = os.path.abspath(brd_file_path)
            abs_dst = os.path.abspath(output_path)
            
            if abs_src != abs_dst:
                shutil.copy2(abs_src, abs_dst)
            
            excel = win32com.client.Dispatch("Excel.Application")
            excel.Visible = False
            excel.DisplayAlerts = False
            
            wb = excel.Workbooks.Open(abs_dst)
            ws = wb.Sheets(sheet_name)
            
            col_num = target_col_index + 1  # 1-based for Excel
            written = 0
            
            for m in matched:
                row_num = m['brd_row']
                try:
                    ws.Cells(row_num, col_num).Value = float(m['new_value'])
                    written += 1
                except (ValueError, TypeError):
                    ws.Cells(row_num, col_num).Value = m['new_value']
                    written += 1
                except Exception as cell_err:
                    logger.warning(f"win32com: Failed to write row {row_num}: {cell_err}")
            
            wb.Save()
            logger.info(f"win32com: Saved {written}/{len(matched)} values to '{abs_dst}'")
            
            wb.Close(False)
            wb = None
            excel.Quit()
            excel = None
            pythoncom.CoUninitialize()
            
            return {
                'matched': matched,
                'unmatched': unmatched,
                'total_written': written,
                'total_skipped': len(unmatched),
                'output_file': output_path,
                'dry_run': False,
            }
            
        except Exception as e:
            logger.error(f"win32com write failed: {e}, falling back to openpyxl")
            if wb:
                try: wb.Close(False)
                except: pass
            if excel:
                try: excel.Quit()
                except: pass
            try: pythoncom.CoUninitialize()
            except: pass
            
            return self._write_matched_with_openpyxl(
                brd_file_path, sheet_name, target_col_index,
                matched, unmatched, output_path
            )

    def _write_matched_with_openpyxl(self, brd_file_path, sheet_name, target_col_index,
                                      matched, unmatched, output_path):
        """Write pre-matched results using openpyxl — fallback."""
        import openpyxl
        
        wb = openpyxl.load_workbook(brd_file_path, keep_links=True)
        ws = wb[sheet_name]
        
        written = 0
        for m in matched:
            row_num = m['brd_row']
            cell = ws.cell(row=row_num, column=target_col_index + 1)
            try:
                cell.value = float(m['new_value'])
                written += 1
            except (ValueError, TypeError):
                cell.value = m['new_value']
                written += 1
        
        wb.save(output_path)
        wb.close()
        logger.info(f"openpyxl: Saved {written}/{len(matched)} values to '{output_path}'")
        
        return {
            'matched': matched,
            'unmatched': unmatched,
            'total_written': written,
            'total_skipped': len(unmatched),
            'output_file': output_path,
            'dry_run': False,
        }
