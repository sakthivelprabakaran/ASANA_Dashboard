"""
Excel Import Dialog - Bulk import test cases from Excel files.
Provides a spreadsheet viewer with column mapping for:
- Test Cases (Name)
- N-Points
- Previous Values
- Perf BRD

Users click column headers to map fields, select a single priority,
and import all rows at once into the Manual Subtask Dialog.
"""

import os
import logging
import openpyxl
from PyQt5.QtWidgets import (QDialog, QVBoxLayout, QHBoxLayout, QLabel, QPushButton,
                             QLineEdit, QTableWidget, QTableWidgetItem, QHeaderView,
                             QMessageBox, QAbstractItemView, QFrame, QComboBox,
                             QFileDialog, QWidget, QStyledItemDelegate, QStyleOptionViewItem,
                             QSizePolicy)
from PyQt5.QtCore import Qt, QModelIndex
from PyQt5.QtGui import QColor, QBrush, QPainter, QFont
from typing import List, Any, Dict, Optional

logger = logging.getLogger('AsanaGenerator.ExcelImportDialog')


# ────────────────────────────────────────────
# Column highlight delegate (reuses BrdViewer pattern)
# ────────────────────────────────────────────

class ImportColumnDelegate(QStyledItemDelegate):
    """Custom delegate for painting column backgrounds based on mapping."""

    def __init__(self, parent=None):
        super().__init__(parent)
        self.col_colors = {}  # col_index -> QColor
        self.COLOR_EVEN = QColor(255, 255, 255)
        self.COLOR_ODD = QColor(248, 250, 252)

    def set_column_colors(self, col_colors: Dict[int, QColor]):
        self.col_colors = col_colors

    def paint(self, painter: QPainter, option: QStyleOptionViewItem, index: QModelIndex):
        table_col = index.column()
        adjusted_col = table_col - 1 if table_col > 0 else -1  # skip row-number column
        row = index.row()

        if adjusted_col in self.col_colors:
            bg = self.col_colors[adjusted_col]
        else:
            bg = self.COLOR_ODD if row % 2 == 1 else self.COLOR_EVEN

        painter.fillRect(option.rect, bg)
        super().paint(painter, option, index)


# ────────────────────────────────────────────
# Main dialog
# ────────────────────────────────────────────

# Mapping field definitions
FIELD_DEFS = [
    {"key": "test_cases",      "label": "🎯 Test Cases",      "color": QColor(134, 239, 172), "required": True},
    {"key": "n_points",        "label": "📊 N-Points",        "color": QColor(147, 197, 253), "required": False},
    {"key": "previous_value",  "label": "📈 Previous Values",  "color": QColor(253, 186, 116), "required": False},
    {"key": "perf_brd",        "label": "📋 Perf BRD",        "color": QColor(196, 181, 253), "required": False},
]


class ExcelImportDialog(QDialog):
    """
    Dialog for bulk-importing test cases from an Excel file.
    Workflow:
      1. Select Excel file & sheet
      2. Click column headers to map fields
      3. Choose a single priority for all rows
      4. Import
    """

    def __init__(self, parent=None):
        super().__init__(parent)
        self.setWindowTitle("📥 Bulk Import from Excel")
        self.setMinimumSize(1200, 750)
        self.setModal(True)

        # Data
        self._raw_sheets: Dict[str, List[List[Any]]] = {}
        self._headers: List[str] = []
        self._data_rows: List[List[Any]] = []

        # Column mapping  field_key -> col_index
        self._mapping: Dict[str, int] = {}
        self._active_field: Optional[str] = None  # which field the next header-click maps

        # Results
        self._imported_subtasks: List[Dict] = []

        self._init_ui()

    # ──────────────────────── UI ────────────────────────

    def _init_ui(self):
        root = QVBoxLayout(self)
        root.setContentsMargins(20, 16, 20, 16)
        root.setSpacing(12)

        # ── Row 1: File picker ──
        file_frame = QFrame()
        file_frame.setStyleSheet("background:#f8fafc; border:1px solid #e2e8f0; border-radius:8px;")
        file_lay = QHBoxLayout(file_frame)
        file_lay.setContentsMargins(14, 10, 14, 10)
        file_lay.setSpacing(10)

        file_lay.addWidget(QLabel("📂 Excel File:"))
        self.lbl_file = QLabel("No file selected")
        self.lbl_file.setStyleSheet("color:#94a3b8; font-size:12px;")
        file_lay.addWidget(self.lbl_file, 1)

        btn_browse = QPushButton("Browse…")
        btn_browse.setCursor(Qt.PointingHandCursor)
        btn_browse.setStyleSheet("padding:8px 16px; background:#3b82f6; color:white; border-radius:6px; font-weight:bold;")
        btn_browse.clicked.connect(self._browse_file)
        file_lay.addWidget(btn_browse)

        file_lay.addWidget(QLabel("Sheet:"))
        self.combo_sheet = QComboBox()
        self.combo_sheet.setMinimumWidth(180)
        self.combo_sheet.currentIndexChanged.connect(self._on_sheet_changed)
        file_lay.addWidget(self.combo_sheet)

        root.addWidget(file_frame)

        # ── Row 2: Instructions + Priority ──
        instr_row = QHBoxLayout()
        instr = QLabel("💡 Click a mapping button below, then click the matching column header in the table")
        instr.setStyleSheet("color:#475569; font-size:12px; padding:6px 10px; background:#f1f5f9; border-radius:6px; border:1px solid #e2e8f0;")
        instr_row.addWidget(instr, 1)

        instr_row.addSpacing(12)
        instr_row.addWidget(QLabel("Priority for all:"))
        self.combo_priority = QComboBox()
        self.combo_priority.addItems(["", "P0", "P1", "P2", "P3"])
        self.combo_priority.setMinimumWidth(90)
        self.combo_priority.setStyleSheet("padding:6px 10px; font-weight:bold;")
        instr_row.addWidget(self.combo_priority)

        root.addLayout(instr_row)

        # ── Row 3: Mapping buttons ──
        map_frame = QFrame()
        map_frame.setStyleSheet("background:#ffffff; border:1px solid #e2e8f0; border-radius:8px;")
        map_lay = QHBoxLayout(map_frame)
        map_lay.setContentsMargins(12, 8, 12, 8)
        map_lay.setSpacing(10)

        map_lay.addWidget(QLabel("Map columns:"))

        self._field_buttons: Dict[str, QPushButton] = {}
        self._field_labels: Dict[str, QLabel] = {}

        for fdef in FIELD_DEFS:
            btn = QPushButton(fdef["label"])
            btn.setCheckable(True)
            btn.setCursor(Qt.PointingHandCursor)
            r, g, b = fdef["color"].red(), fdef["color"].green(), fdef["color"].blue()
            required_tag = " *" if fdef["required"] else ""
            btn.setText(fdef["label"] + required_tag)
            btn.setStyleSheet(f"""
                QPushButton {{
                    padding:8px 14px; border-radius:6px; font-weight:bold; font-size:12px;
                    background:#f1f5f9; color:#475569; border:2px solid #e2e8f0;
                }}
                QPushButton:checked {{
                    background:rgba({r},{g},{b},60); border-color:rgba({r},{g},{b},200); color:#1e293b;
                }}
            """)
            btn.clicked.connect(lambda checked, k=fdef["key"]: self._activate_field(k))
            map_lay.addWidget(btn)
            self._field_buttons[fdef["key"]] = btn

            lbl = QLabel("—")
            lbl.setStyleSheet("color:#94a3b8; font-size:11px; min-width:80px;")
            map_lay.addWidget(lbl)
            self._field_labels[fdef["key"]] = lbl

        map_lay.addStretch()

        # Reset button
        btn_reset = QPushButton("↺ Reset")
        btn_reset.setCursor(Qt.PointingHandCursor)
        btn_reset.setStyleSheet("padding:6px 12px; background:#fee2e2; color:#dc2626; border:1px solid #fecaca; border-radius:6px; font-weight:bold;")
        btn_reset.clicked.connect(self._reset_mapping)
        map_lay.addWidget(btn_reset)

        root.addWidget(map_frame)

        # ── Row 4: Spreadsheet table ──
        self.table = QTableWidget()
        self.table.setAlternatingRowColors(False)
        self.table.setSelectionMode(QAbstractItemView.NoSelection)
        self.table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.Interactive)
        self.table.horizontalHeader().setMinimumSectionSize(60)
        self.table.horizontalHeader().setDefaultSectionSize(130)
        self.table.verticalHeader().setVisible(False)
        self.table.horizontalHeader().sectionClicked.connect(self._on_header_clicked)
        self.table.horizontalHeader().setStyleSheet("""
            QHeaderView::section {
                background-color:#f1f5f9; color:#475569; padding:8px 4px;
                border:1px solid #e2e8f0; font-weight:bold; font-size:11px;
            }
            QHeaderView::section:hover {
                background-color:#e2e8f0; color:#1e40af;
            }
        """)
        self.table.setStyleSheet("""
            QTableWidget {
                background:#ffffff; border:1px solid #e2e8f0; border-radius:8px;
                gridline-color:#f1f5f9; font-size:12px;
            }
            QTableWidget::item { padding:4px 6px; }
        """)

        self.delegate = ImportColumnDelegate(self.table)
        self.table.setItemDelegate(self.delegate)

        root.addWidget(self.table, 1)

        # ── Row 5: Summary + action buttons ──
        bottom = QHBoxLayout()
        bottom.setSpacing(12)

        self.lbl_summary = QLabel("Import an Excel file to begin")
        self.lbl_summary.setStyleSheet("color:#64748b; font-size:13px; font-weight:bold;")
        bottom.addWidget(self.lbl_summary, 1)

        btn_cancel = QPushButton("Cancel")
        btn_cancel.setCursor(Qt.PointingHandCursor)
        btn_cancel.setStyleSheet("padding:10px 20px; background:#f1f5f9; color:#475569; border:1px solid #e2e8f0; border-radius:6px; font-weight:bold;")
        btn_cancel.clicked.connect(self.reject)
        bottom.addWidget(btn_cancel)

        self.btn_import = QPushButton("✅ Import All")
        self.btn_import.setCursor(Qt.PointingHandCursor)
        self.btn_import.setStyleSheet("padding:10px 24px; background:#10b981; color:white; border:none; border-radius:6px; font-weight:bold; font-size:13px;")
        self.btn_import.clicked.connect(self._do_import)
        self.btn_import.setEnabled(False)
        bottom.addWidget(self.btn_import)

        root.addLayout(bottom)

    # ──────────────────────── File loading ────────────────────────

    def _browse_file(self):
        fname, _ = QFileDialog.getOpenFileName(
            self, "Select Excel File", "",
            "Excel Files (*.xlsx *.xls);;All Files (*)"
        )
        if not fname:
            return

        try:
            wb = openpyxl.load_workbook(fname, data_only=True, read_only=True)
            self._raw_sheets = {}
            for sname in wb.sheetnames:
                ws = wb[sname]
                rows = []
                for row in ws.iter_rows():
                    rows.append([cell.value for cell in row])
                self._raw_sheets[sname] = rows
            wb.close()

            self.lbl_file.setText(f"✓ {os.path.basename(fname)}")
            self.lbl_file.setStyleSheet("color:#10b981; font-size:12px; font-weight:bold;")

            self.combo_sheet.blockSignals(True)
            self.combo_sheet.clear()
            for sname, data in self._raw_sheets.items():
                row_count = max(0, len(data) - 1)
                self.combo_sheet.addItem(f"{sname} ({row_count} rows)", sname)
            self.combo_sheet.blockSignals(False)

            self._on_sheet_changed()

        except Exception as e:
            logger.error(f"Failed to load Excel: {e}")
            QMessageBox.critical(self, "Error", f"Failed to load file:\n{e}")

    def _on_sheet_changed(self):
        sname = self.combo_sheet.currentData()
        if not sname or sname not in self._raw_sheets:
            return

        raw = self._raw_sheets[sname]
        if not raw:
            return

        self._headers = [
            str(h).strip() if h and str(h).lower() not in ('nan', 'none', '') else f"Col {i}"
            for i, h in enumerate(raw[0])
        ]
        self._data_rows = raw[1:]

        self._reset_mapping()
        self._populate_table()
        self._update_summary()

    def _populate_table(self):
        max_rows = min(len(self._data_rows), 500)
        col_count = len(self._headers)

        self.table.setRowCount(max_rows)
        self.table.setColumnCount(col_count + 1)  # +1 for row-number column

        header_labels = ["#"] + [self._trunc(h, 40) for h in self._headers]
        self.table.setHorizontalHeaderLabels(header_labels)

        for r_idx in range(max_rows):
            row_data = self._data_rows[r_idx] if r_idx < len(self._data_rows) else []

            # Row number
            num_item = QTableWidgetItem(str(r_idx + 1))
            num_item.setTextAlignment(Qt.AlignCenter)
            num_item.setForeground(QBrush(QColor("#64748b")))
            self.table.setItem(r_idx, 0, num_item)

            for c_idx in range(col_count):
                val = row_data[c_idx] if c_idx < len(row_data) else None
                display = str(val).strip() if val is not None else ""
                if display.lower() in ('nan', 'none', 'null', ''):
                    display = ""

                item = QTableWidgetItem(self._trunc(display, 50))
                item.setToolTip(str(val) if val is not None else "Empty")
                if not display:
                    item.setForeground(QBrush(QColor("#94a3b8")))
                    item.setText("—")
                self.table.setItem(r_idx, c_idx + 1, item)

        self._refresh_highlighting()

    # ──────────────────────── Column mapping ────────────────────────

    def _activate_field(self, field_key: str):
        """Activate a mapping field — next header click maps to this field."""
        # Uncheck all other buttons
        for k, btn in self._field_buttons.items():
            btn.setChecked(k == field_key)
        self._active_field = field_key

    def _on_header_clicked(self, logical_index: int):
        if logical_index == 0:  # skip row-number column
            return
        if not self._active_field:
            QMessageBox.information(self, "Select a field first",
                "Click one of the mapping buttons (Test Cases, N-Points, etc.) first,\n"
                "then click the column header.")
            return

        col_index = logical_index - 1
        col_name = self._headers[col_index] if col_index < len(self._headers) else f"Col {col_index}"

        # Remove this column from any other mapping
        for k in list(self._mapping.keys()):
            if self._mapping[k] == col_index:
                del self._mapping[k]
                self._field_labels[k].setText("—")
                self._field_labels[k].setStyleSheet("color:#94a3b8; font-size:11px; min-width:80px;")

        # Set mapping
        self._mapping[self._active_field] = col_index
        self._field_labels[self._active_field].setText(f"✓ {col_name}")
        self._field_labels[self._active_field].setStyleSheet("color:#059669; font-size:11px; font-weight:bold; min-width:80px;")

        # Auto-advance to next unmapped field
        self._auto_advance()
        self._refresh_highlighting()
        self._update_summary()

    def _auto_advance(self):
        """Auto-advance active field to the next unmapped field."""
        order = [f["key"] for f in FIELD_DEFS]
        for key in order:
            if key not in self._mapping:
                self._activate_field(key)
                return
        # All mapped — deactivate
        self._active_field = None
        for btn in self._field_buttons.values():
            btn.setChecked(False)

    def _reset_mapping(self):
        self._mapping = {}
        self._active_field = None
        for k in self._field_buttons:
            self._field_buttons[k].setChecked(False)
            self._field_labels[k].setText("—")
            self._field_labels[k].setStyleSheet("color:#94a3b8; font-size:11px; min-width:80px;")
        # Auto-activate first field
        if FIELD_DEFS:
            self._activate_field(FIELD_DEFS[0]["key"])
        self._refresh_highlighting()
        self._update_summary()

    def _refresh_highlighting(self):
        """Update delegate colors from current mapping."""
        col_colors: Dict[int, QColor] = {}
        for fdef in FIELD_DEFS:
            idx = self._mapping.get(fdef["key"], -1)
            if idx >= 0:
                col_colors[idx] = fdef["color"]
        self.delegate.set_column_colors(col_colors)
        self.table.viewport().update()

    def _update_summary(self):
        tc_col = self._mapping.get("test_cases", -1)
        if tc_col < 0:
            self.lbl_summary.setText("⚠️ Map the 'Test Cases' column (required)")
            self.btn_import.setEnabled(False)
            return

        # Count non-empty rows in the Test Cases column
        count = 0
        for row in self._data_rows:
            val = row[tc_col] if tc_col < len(row) else None
            if val and str(val).strip() and str(val).strip().lower() not in ('nan', 'none'):
                count += 1

        mapped_fields = [f["label"] for f in FIELD_DEFS if f["key"] in self._mapping]
        priority = self.combo_priority.currentText() or "none"
        self.lbl_summary.setText(
            f"✅ Ready to import {count} test cases  |  Priority: {priority}  |  Mapped: {', '.join(mapped_fields)}"
        )
        self.btn_import.setEnabled(count > 0)

    # ──────────────────────── Import ────────────────────────

    def _do_import(self):
        tc_col = self._mapping.get("test_cases", -1)
        if tc_col < 0:
            QMessageBox.warning(self, "Missing Mapping", "Please map the 'Test Cases' column first.")
            return

        np_col = self._mapping.get("n_points", -1)
        pv_col = self._mapping.get("previous_value", -1)
        pb_col = self._mapping.get("perf_brd", -1)
        priority = self.combo_priority.currentText()

        results = []
        for row in self._data_rows:
            name_val = row[tc_col] if tc_col < len(row) else None
            name = str(name_val).strip() if name_val else ""
            if not name or name.lower() in ('nan', 'none'):
                continue

            n_points = self._cell(row, np_col)
            prev_val = self._cell(row, pv_col)
            perf_brd = self._cell(row, pb_col)

            results.append({
                "name": name,
                "priority": priority,
                "est_time": "",
                "n_points": n_points,
                "previous_value": prev_val,
                "perf_brd": perf_brd,
            })

        self._imported_subtasks = results
        logger.info(f"Excel import: {len(results)} subtasks imported")
        self.accept()

    def get_imported_subtasks(self) -> List[Dict]:
        return list(self._imported_subtasks)

    # ──────────────────────── Helpers ────────────────────────

    @staticmethod
    def _cell(row: list, col_idx: int) -> str:
        if col_idx < 0 or col_idx >= len(row):
            return ""
        val = row[col_idx]
        if val is None:
            return ""
        s = str(val).strip()
        return "" if s.lower() in ('nan', 'none') else s

    @staticmethod
    def _trunc(text: str, length: int) -> str:
        return text[:length] + "…" if len(text) > length else text